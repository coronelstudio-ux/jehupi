"""Resumen en Excel para cargar compras y ventas en Contadis (Editora ADIS) a mano, lo más rápido posible.

Contadis no tiene importación: el archivo sigue el orden de su pantalla "Entrada de datos"
(Factura en 3 casillas, Comprobante, Razón social, Fecha, Condición, Timbrado, RUC) y del panel
de montos (Gravadas 5%, Gravadas 10%, I.V.A. 5%, I.V.A. 10%, Exenta, Total), donde las gravadas
van SIN IVA. La app guarda los montos con IVA incluido; aquí se separan base e IVA.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .cuentas import descripcion_contadis, etiqueta
from .validaciones import NUMERO_SIN_NOMBRE, TOLERANCIA_GS, parsear_fecha

NOMBRE_SIN_NOMBRE = "CLIENTE OCASIONAL"  # como aparece en Contadis para el consumidor final
COMPROBANTE = {109: "FACTURA", 110: "NOTA DE CREDITO", 111: "NOTA DE DEBITO", 112: "TICKET"}
CONDICION = {1: "Contado", 2: "Crédito"}
ESTADOS_INCLUIDOS = ("OK", "APROBADA", "EXCLUIDA")

MONTOS = ["Gravadas 5%", "Gravadas 10%", "I.V.A. 5%", "I.V.A. 10%", "Exenta", "Total"]
# Mismo orden que la pantalla "Entrada de datos" de Contadis; después, el panel de montos.
COLUMNAS = {
    "COMPRAS": ["Est.", "Pto.", "Número", "Comprobante", "Razón social", "Fecha", "Condición", "Cuenta",
                "Descripción", "Timbrado", "RUC", *MONTOS, "Concepto", "Detalle", "Electrónica", "Cargado",
                "Observaciones"],
    "VENTAS": ["Est.", "Pto.", "Número", "Comprobante", "Razón social", "RUC / CI", "Fecha", "Condición",
               "Cuenta", "Descripción", "Timbrado", *MONTOS, "Concepto", "Detalle", "Electrónica", "Cargado",
               "Observaciones"],
}

AZUL, VERDE, AMARILLO = "1F4E78", "C6EFCE", "FFF2CC"
FORMATO_GS = "#,##0"


def separar_iva(gravada_con_iva: int, iva_impreso: int, divisor: int) -> tuple[int, int]:
    """(base sin IVA, IVA). Usa el IVA impreso si coincide con el calculado; si no, el calculado.

    Divisor 11 para el 10% y 21 para el 5%: base + IVA = gravada con IVA, sin diferencias de redondeo.
    """
    gravada_con_iva = int(gravada_con_iva or 0)
    calculado = round(gravada_con_iva / divisor)
    iva = int(iva_impreso) if iva_impreso and abs(int(iva_impreso) - calculado) <= TOLERANCIA_GS else calculado
    return gravada_con_iva - iva, iva


def _partes_numero(numero: str) -> tuple[str, str, str]:
    partes = str(numero or "").split("-")
    return tuple(partes) if len(partes) == 3 else ("", "", str(numero or ""))


def _documento(c: dict, tipo_lote: str) -> str:
    numero, dv = str(c.get("ruc") or ""), str(c.get("dv") or "")
    if tipo_lote == "VENTAS" and str(c.get("tipo_identificacion") or "11") != "11":
        return numero
    return f"{numero}-{dv}" if numero and dv else numero


def _razon_social(c: dict, tipo_lote: str) -> str:
    if tipo_lote == "VENTAS" and (str(c.get("tipo_identificacion")) == "15" or c.get("ruc") == NUMERO_SIN_NOMBRE):
        return NOMBRE_SIN_NOMBRE
    return (c.get("razon_social") or "").strip()


def fila(c: dict, tipo_lote: str, origen: str = "", cuentas_por_id: dict | None = None) -> tuple[list, list[str]]:
    """(valores en el orden de Contadis, avisos para la columna Observaciones)."""
    cuenta = (cuentas_por_id or {}).get(c.get("cuenta_id"))
    base5, iva5 = separar_iva(c.get("gravada_5"), c.get("iva_5"), 21)
    base10, iva10 = separar_iva(c.get("gravada_10"), c.get("iva_10"), 11)
    exenta = int(c.get("exenta") or 0)
    total = base5 + base10 + iva5 + iva10 + exenta
    avisos = []
    if c.get("total") and abs(total - int(c["total"])) > TOLERANCIA_GS:
        avisos.append(f"Total de la factura: {int(c['total']):,}".replace(",", "."))
    if c["estado"] == "APROBADA" and c.get("errores"):
        avisos.append("Aprobada con observaciones")
    razon = _razon_social(c, tipo_lote)
    if not razon:
        avisos.append("Falta razón social")
    if not cuenta:
        avisos.append("Sin cuenta")
    elif c.get("cuenta_origen") == "IA" and c["estado"] != "APROBADA":  # aprobada = confirmada por una persona
        avisos.append("Cuenta sugerida por la IA")
    est, pto, nro = _partes_numero(c.get("numero"))
    fecha = parsear_fecha(c.get("fecha"))
    cabecera = [est, pto, nro, COMPROBANTE.get(c.get("tipo_comprobante"), str(c.get("tipo_comprobante") or ""))]
    contraparte = [razon]
    if tipo_lote == "VENTAS":
        contraparte.append(_documento(c, tipo_lote))
    datos = [fecha, CONDICION.get(c.get("condicion"), ""), etiqueta(cuenta),
             descripcion_contadis(c, cuenta, razon) if cuenta else "", str(c.get("timbrado") or "")]
    if tipo_lote == "COMPRAS":
        datos.append(_documento(c, tipo_lote))
    montos = [base5, base10, iva5, iva10, exenta, total]
    resto = [c.get("concepto") or "", c.get("detalle") or "", "Sí" if c.get("clase") == "ELECTRONICA" else "", "",
             "; ".join(avisos + ([origen] if origen else []))]
    return cabecera + contraparte + datos + montos + resto, avisos


def _ordenar(comprobantes: list[dict]) -> list[dict]:
    return sorted(comprobantes, key=lambda c: (parsear_fecha(c.get("fecha")) or datetime.max.date(),
                                               str(c.get("numero") or "")))


def separar(comprobantes: list[dict]) -> tuple[list[dict], list[dict]]:
    """(para cargar, pendientes). Pendientes: sin revisar, o electrónicas a las que les faltan los montos."""
    cargar, pendientes = [], []
    for c in comprobantes:
        if c["estado"] == "DESCARTADA":
            continue
        if c["estado"] in ESTADOS_INCLUIDOS and int(c.get("total") or 0) > 0:
            cargar.append(c)
        else:
            pendientes.append(c)
    return _ordenar(cargar), _ordenar(pendientes)


def _encabezado(ws, columnas: list[str]) -> None:
    ws.append(columnas)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def _hoja_registro(wb: Workbook, tipo_lote: str, comprobantes: list[dict], paginas: dict,
                   cuentas_por_id: dict | None = None) -> dict:
    ws = wb.create_sheet(tipo_lote.capitalize())
    columnas = COLUMNAS[tipo_lote]
    _encabezado(ws, columnas)
    for c in comprobantes:
        valores, _ = fila(c, tipo_lote, paginas.get(c.get("pagina_id"), ""), cuentas_por_id)
        ws.append(valores)
    n = len(comprobantes)
    col = {nombre: get_column_letter(i) for i, nombre in enumerate(columnas, start=1)}

    # Formatos: fecha dd/mm/aaaa, montos con separador de miles, textos que parecen números como texto.
    for fila_celdas in ws.iter_rows(min_row=2, max_row=n + 1):
        for celda in fila_celdas:
            nombre = columnas[celda.column - 1]
            if nombre == "Fecha":
                celda.number_format = "DD/MM/YYYY"
            elif nombre in MONTOS:
                celda.number_format = FORMATO_GS
            elif nombre in ("Est.", "Pto.", "Número", "Timbrado", "RUC", "RUC / CI"):
                celda.number_format = "@"

    # Totales al pie, para comparar con los totales de Contadis.
    totales = {}
    if n:
        ws.append([])
        fila_total = n + 3
        ws.cell(fila_total, 1, "TOTALES").font = Font(bold=True)
        for nombre in MONTOS:
            letra = col[nombre]
            celda = ws[f"{letra}{fila_total}"]
            celda.value = f"=SUM({letra}2:{letra}{n + 1})"
            celda.number_format = FORMATO_GS
            celda.font = Font(bold=True)
        for c in comprobantes:
            valores, _ = fila(c, tipo_lote)
            for nombre in MONTOS:
                totales[nombre] = totales.get(nombre, 0) + valores[columnas.index(nombre)]

        # Columna "Cargado": lista Sí/No y la fila se pinta de verde al marcarla.
        lista = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
        ws.add_data_validation(lista)
        lista.add(f"{col['Cargado']}2:{col['Cargado']}{n + 1}")
        ultima = get_column_letter(len(columnas))
        ws.conditional_formatting.add(
            f"A2:{ultima}{n + 1}",
            FormulaRule(formula=[f'${col["Cargado"]}2="Sí"'], fill=PatternFill("solid", fgColor=VERDE)))
        ws.auto_filter.ref = f"A1:{ultima}{n + 1}"

    anchos = {"Est.": 6, "Pto.": 6, "Número": 10, "Comprobante": 13, "Razón social": 34, "Fecha": 11,
              "Condición": 10, "Cuenta": 34, "Descripción": 50, "Timbrado": 11, "RUC": 13, "RUC / CI": 13,
              "Concepto": 28, "Detalle": 45, "Electrónica": 11, "Cargado": 9, "Observaciones": 40}
    for nombre, letra in col.items():
        ws.column_dimensions[letra].width = anchos.get(nombre, 13)
    return {"cantidad": n, **totales}


def generar_resumen(compras: list[dict], ventas: list[dict], cliente: dict, periodo: str,
                    paginas: dict[int, str] | None = None, cuentas_por_id: dict | None = None) -> bytes:
    """Excel con hojas Compras, Ventas, Pendientes y Resumen, listo para cargar en Contadis."""
    paginas = paginas or {}
    wb = Workbook()
    wb.remove(wb.active)
    compras_ok, compras_pend = separar(compras)
    ventas_ok, ventas_pend = separar(ventas)
    totales = {
        "COMPRAS": _hoja_registro(wb, "COMPRAS", compras_ok, paginas, cuentas_por_id),
        "VENTAS": _hoja_registro(wb, "VENTAS", ventas_ok, paginas, cuentas_por_id),
    }

    ws = wb.create_sheet("Pendientes")
    _encabezado(ws, ["Registro", "Estado", "Número", "Fecha", "RUC / CI", "Razón social", "Total", "Motivo", "Origen"])
    for tipo, lista in (("COMPRAS", compras_pend), ("VENTAS", ventas_pend)):
        for c in lista:
            motivo = "; ".join(c.get("errores") or []) or (
                "Electrónica sin montos: completar en Revisión" if c.get("clase") == "ELECTRONICA" else "Sin revisar")
            ws.append([tipo, c["estado"], c.get("numero") or "", c.get("fecha") or "", _documento(c, tipo),
                       c.get("razon_social") or "", int(c.get("total") or 0), motivo,
                       paginas.get(c.get("pagina_id"), "")])
    for letra, ancho in zip("ABCDEFGHI", (10, 11, 17, 11, 13, 30, 13, 60, 22)):
        ws.column_dimensions[letra].width = ancho
    for celda in ws["G"][1:]:
        celda.number_format = FORMATO_GS

    ws = wb.create_sheet("Resumen", 0)
    ws.append(["Resumen para cargar en Contadis"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Cliente: {cliente['nombre']} — RUC {cliente['ruc']}-{cliente['dv']}"])
    ws.append([f"Período: {periodo}    Generado: {datetime.now():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append(["Registro", "Para cargar", "Pendientes", *MONTOS])
    for celda in ws[5]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
    for tipo, pend in (("COMPRAS", compras_pend), ("VENTAS", ventas_pend)):
        t = totales[tipo]
        ws.append([tipo.capitalize(), t["cantidad"], len(pend), *(t.get(m, 0) for m in MONTOS)])
    for fila_celdas in ws.iter_rows(min_row=6, max_row=7, min_col=4, max_col=9):
        for celda in fila_celdas:
            celda.number_format = FORMATO_GS
    ws.append([])
    notas = [
        "Gravadas 5% y 10% van SIN IVA, como en el panel de montos de Contadis; el IVA va aparte.",
        "Las facturas electrónicas se cargan en Contadis aunque no se importen a Marangatu (columna Electrónica).",
        "Cuenta: la recordada para ese proveedor o la sugerida por la IA (ver Observaciones). Revísela al cargar.",
        f"Consumidor final: razón social \"{NOMBRE_SIN_NOMBRE}\".",
        "Marque la columna Cargado con \"Sí\" a medida que avanza: la fila se pinta de verde.",
        "Al terminar, compare los totales del pie de cada hoja con los de Contadis.",
    ]
    for nota in notas:
        ws.append([nota])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=AMARILLO)
    ws.column_dimensions["A"].width = 16
    for letra in "BCDEFGHI":
        ws.column_dimensions[letra].width = 14

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
