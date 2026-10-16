"""Generación del archivo de importación de Marangatu (RG 90/2021) y del Excel de control."""
import io
import json
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .validaciones import IDENT_SIN_RAZON_SOCIAL, TIPOS_SOPORTADOS

COLUMNAS_COMPRAS = [
    "Tipo de registro", "Tipo de identificación", "Número de identificación", "Razón social",
    "Tipo de comprobante", "Fecha de emisión", "Timbrado", "Número de comprobante",
    "Gravada 10% (IVA incl.)", "Gravada 5% (IVA incl.)", "Exenta", "Total",
    "Condición", "Moneda extranjera", "Imputa IVA", "Imputa IRE", "Imputa IRP-RSP", "No imputa",
    "Nro. comprobante asociado", "Timbrado comprobante asociado",
]

COLUMNAS_VENTAS = [c for c in COLUMNAS_COMPRAS if c != "No imputa"]  # 19 campos

# Letra del identificador del archivo: distinta para compras y ventas, para que no se repitan nombres.
LETRA_LOTE = {"COMPRAS": "C", "VENTAS": "V"}

FIN_DE_LINEA = "\r\n"


def _sn(valor) -> str:
    return "S" if valor else "N"


def fila_compra(c: dict) -> list[str]:
    """Una fila de 20 campos del registro de compras."""
    es_ticket = c["tipo_comprobante"] == 112
    g10, g5, ex = (0, 0, 0) if es_ticket else (c["gravada_10"] or 0, c["gravada_5"] or 0, c["exenta"] or 0)
    total = c["total"] if es_ticket else g10 + g5 + ex
    imputa = [bool(c["imputa_iva"]), bool(c["imputa_ire"]), bool(c["imputa_irp"])]
    return [
        "2",                        # 1  tipo de registro: compras
        "11",                       # 2  tipo de identificación: RUC
        str(c["ruc"]),              # 3  RUC sin DV
        "",                         # 4  razón social: la completa Marangatu por el RUC
        str(c["tipo_comprobante"]),  # 5  109 factura / 112 ticket
        c["fecha"],                 # 6  dd/mm/aaaa
        str(c["timbrado"]),         # 7
        c["numero"],                # 8  001-001-0000000
        str(g10), str(g5), str(ex), str(total),  # 9-12
        str(c["condicion"]),        # 13 1 contado / 2 crédito
        "N",                        # 14 moneda extranjera
        *(_sn(v) for v in imputa),  # 15-17
        _sn(not any(imputa)),       # 18 no imputa
        "", "",                     # 19-20 solo notas de crédito/débito
    ]


def fila_venta(c: dict) -> list[str]:
    """Una fila de 19 campos del registro de ventas (no tiene la columna "no imputa")."""
    g10, g5, ex = c["gravada_10"] or 0, c["gravada_5"] or 0, c["exenta"] or 0
    total = g10 + g5 + ex if c["tipo_comprobante"] == 109 else c["total"]
    tipo_ident = str(c.get("tipo_identificacion") or "11")
    razon_social = "" if tipo_ident in IDENT_SIN_RAZON_SOCIAL else (c.get("razon_social") or "").strip()
    imputa = [bool(c["imputa_iva"]), bool(c["imputa_ire"]), bool(c["imputa_irp"])]
    return [
        "1",                        # 1  tipo de registro: ventas
        tipo_ident,                 # 2  11 RUC, 12 cédula, 15 sin nombre…
        str(c["ruc"]),              # 3  RUC sin DV, cédula sin puntos o 44444401
        razon_social,               # 4  en blanco para RUC, cédula y sin nombre
        str(c["tipo_comprobante"]),  # 5
        c["fecha"],                 # 6
        str(c["timbrado"]),         # 7  timbrado propio del cliente
        c["numero"],                # 8
        str(g10), str(g5), str(ex), str(total),  # 9-12 (en ventas también para tickets)
        str(c["condicion"]),        # 13
        "N",                        # 14 moneda extranjera
        *(_sn(v) for v in imputa),  # 15-17
        "", "",                     # 18-19 solo notas de crédito/débito
    ]


def fila(c: dict, tipo_lote: str) -> list[str]:
    return fila_venta(c) if tipo_lote == "VENTAS" else fila_compra(c)


def exportables(comprobantes: list[dict]) -> tuple[list[dict], list[dict]]:
    """Separa los que van al archivo de los que no (pendientes, electrónicos, tipos no soportados)."""
    van, no_van = [], []
    for c in comprobantes:
        if c["estado"] in ("OK", "APROBADA") and c["clase"] != "ELECTRONICA" and c["tipo_comprobante"] in TIPOS_SOPORTADOS:
            van.append(c)
        else:
            no_van.append(c)
    return van, no_van


def nombre_archivo(ruc: str, periodo: str, secuencia: int, tipo_lote: str = "COMPRAS") -> str:
    mes, anio = periodo.split("/")
    return f"{ruc}_REG_{int(mes):02d}{anio}_{LETRA_LOTE[tipo_lote]}{secuencia:04d}"


def generar_zips(comprobantes: list[dict], ruc: str, periodo: str, primera_secuencia: int,
                 tipo_lote: str = "COMPRAS") -> list[tuple[str, bytes, int]]:
    """Devuelve [(nombre_zip, contenido, cantidad)]. Divide en varios archivos si supera el máximo de filas."""
    filas = [fila(c, tipo_lote) for c in comprobantes]
    salida = []
    for i in range(0, len(filas), config.MAX_FILAS_ARCHIVO):
        bloque = filas[i:i + config.MAX_FILAS_ARCHIVO]
        nombre = nombre_archivo(ruc, periodo, primera_secuencia + len(salida), tipo_lote)
        texto = FIN_DE_LINEA.join("\t".join(f) for f in bloque) + FIN_DE_LINEA
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr(f"{nombre}.txt", texto.encode("utf-8"))
        salida.append((f"{nombre}.zip", buffer.getvalue(), len(bloque)))
    return salida


def _hoja(wb: Workbook, titulo: str, encabezados: list[str], filas: list[list]) -> None:
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F4E78")
    for f in filas:
        ws.append(f)
    for i, enc in enumerate(encabezados, start=1):
        ancho = max([len(str(enc))] + [len(str(f[i - 1])) for f in filas[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(ancho + 2, 60)
    ws.freeze_panes = "A2"


def generar_excel_control(van: list[dict], no_van: list[dict], paginas: dict[int, str],
                          tipo_lote: str = "COMPRAS") -> bytes:
    """Excel para revisión interna (con encabezados). No se sube a Marangatu."""
    wb = Workbook()
    wb.remove(wb.active)
    columnas = COLUMNAS_VENTAS if tipo_lote == "VENTAS" else COLUMNAS_COMPRAS
    extra = (["Nombre leído"] if tipo_lote == "VENTAS" else []) + ["Estado", "Clase", "Origen"]
    _hoja(wb, "A importar", columnas + extra,
          [fila(c, tipo_lote) + ([c.get("razon_social") or ""] if tipo_lote == "VENTAS" else [])
           + [c["estado"], c["clase"], paginas.get(c["pagina_id"], "")] for c in van])
    electronicas = [c for c in no_van if c["clase"] == "ELECTRONICA" and c["estado"] == "EXCLUIDA"]
    _hoja(wb, "Electrónicas (control)", ["CDC", "RUC", "DV", "Número", "Fecha", "Origen", "Observaciones"],
          [[c["cdc"], c["ruc"], c["dv"], c["numero"], c["fecha"], paginas.get(c["pagina_id"], ""),
            c["observaciones"]] for c in electronicas])
    otros = [c for c in no_van if c not in electronicas]
    _hoja(wb, "No incluidos", ["ID", "Estado", "Clase", "RUC", "Timbrado", "Número", "Fecha", "Total", "Origen", "Motivo"],
          [[c["id"], c["estado"], c["clase"], c["ruc"], c["timbrado"], c["numero"], c["fecha"], c["total"],
            paginas.get(c["pagina_id"], ""), "; ".join(json.loads(c["errores"] or "[]")
                                                       if isinstance(c["errores"], str) else c["errores"])]
           for c in otros])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
