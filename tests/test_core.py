import io
import json
import zipfile

import cv2
import numpy as np
import pymupdf
import pytest

from core import cdc as cdc_mod
from core import archivos, config, db, exportar, procesador
from core.validaciones import calcular_dv, normalizar_numero, normalizar_ruc, validar_comprobante


def _cdc_valido(base43: str) -> str:
    return base43 + str(calcular_dv(base43))


CDC = _cdc_valido("01" + "80012345".zfill(8) + str(calcular_dv("80012345")) + "001" + "002" + "0000123"
                  + "2" + "20260915" + "1" + "123456789")


# ---------- validaciones ----------

@pytest.mark.parametrize("ruc,dv", [("4327185", 5), ("1895550", 9), ("5850580", 6)])
def test_dv_ruc(ruc, dv):
    assert calcular_dv(ruc) == dv


def test_normalizaciones():
    assert normalizar_ruc("1.895.550-9") == ("1895550", "9")
    assert normalizar_numero("1-1-45645") == "001-001-0045645"
    assert normalizar_numero("0010010045645") == "001-001-0045645"


def _comprobante(**cambios):
    c = {"tipo_comprobante": 109, "ruc": "1895550", "dv": "9", "timbrado": "18586180", "numero": "001-001-0045645",
         "fecha": "09/04/2026", "vigencia_inicio": "19/01/2026", "vigencia_fin": "31/01/2027",
         "gravada_10": 51000, "gravada_5": 0, "exenta": 0, "total": 51000, "iva_10": 4636, "iva_5": 0,
         "condicion": 1, "moneda": "PYG", "ruc_receptor": "5850580", "campos_dudosos": []}
    c.update(cambios)
    return c


def test_comprobante_correcto():
    assert validar_comprobante(_comprobante(), "5850580", "04/2026") == []


@pytest.mark.parametrize("cambio,texto", [
    ({"dv": "3"}, "DV del RUC del proveedor incorrecto"),
    ({"timbrado": "1858618"}, "timbrado"),
    ({"total": 60000}, "no coincide con el total"),
    ({"iva_10": 5000}, "IVA 10%"),
    ({"fecha": "09/05/2026"}, "período"),
    ({"fecha": "09/02/2027"}, "fin de vigencia"),
    ({"ruc_receptor": "123"}, "a nombre del RUC"),
    ({"condicion": None}, "contado o crédito"),
    ({"campos_dudosos": ["total"]}, "dudosos"),
])
def test_comprobante_con_errores(cambio, texto):
    errores = validar_comprobante(_comprobante(**cambio), "5850580", "04/2026")
    assert any(texto in e for e in errores), errores


# ---------- CDC ----------

def test_cdc():
    assert cdc_mod.dv_valido(CDC)
    assert not cdc_mod.dv_valido(CDC[:-1] + str((int(CDC[-1]) + 1) % 10))
    datos = cdc_mod.parsear(CDC)
    assert datos["ruc"] == "80012345" and datos["numero"] == "001-002-0000123" and datos["fecha"] == "15/09/2026"
    agrupado = " ".join(CDC[i:i + 4] for i in range(0, 44, 4))
    assert cdc_mod.buscar_en_texto(f"CDC: {agrupado}\nTotal 100") == [CDC]
    assert cdc_mod.desde_qr(f"https://ekuatia.set.gov.py/consultas/qr?nVersion=150&Id={CDC}&dFeEmiDE=x") == CDC


# ---------- proceso completo con un extractor simulado ----------

def _pdf(paginas_texto: list[str]) -> bytes:
    doc = pymupdf.open()
    for t in paginas_texto:
        doc.new_page().insert_text((72, 72), t)
    return doc.tobytes()


def _png_con_qr(texto: str) -> bytes:
    qr = cv2.QRCodeEncoder.create().encode(texto)
    qr = cv2.resize(qr, None, fx=8, fy=8, interpolation=cv2.INTER_NEAREST)
    lienzo = np.full((qr.shape[0] + 200, qr.shape[1] + 200), 255, np.uint8)
    lienzo[100:-100, 100:-100] = qr
    return cv2.imencode(".png", lienzo)[1].tobytes()


def _item(**cambios):
    item = {"clase": "PREIMPRESA", "tipo_documento": "FACTURA", "cdc": "", "ruc_emisor": "1895550", "dv_emisor": "9",
            "razon_social_emisor": "Librería", "ruc_receptor": "5850580", "dv_receptor": "6", "timbrado": "18586180",
            "vigencia_inicio": "19/01/2026", "vigencia_fin": "31/01/2027", "numero": "001-001-45645",
            "fecha_emision": "09/04/2026", "condicion": "CONTADO", "moneda": "PYG", "gravada_10": 51000,
            "gravada_5": 0, "exenta": 0, "total": 51000, "iva_10": 4636, "iva_5": 0, "campos_dudosos": [],
            "observaciones": ""}
    item.update(cambios)
    return item


def _cliente_y_lote(periodo="04/2026"):
    cliente_id = db.ejecutar("INSERT INTO clientes (ruc, dv, nombre) VALUES ('5850580','6','Cliente Prueba') RETURNING id")
    lote_id = db.ejecutar("INSERT INTO lotes (cliente_id, tipo, periodo, creado_por) "
                          "VALUES (%s, 'COMPRAS', %s, 'test') RETURNING id", (cliente_id, periodo))
    return cliente_id, lote_id


def test_proceso_y_exportacion(base):
    cliente_id, lote_id = _cliente_y_lote()
    llamadas = []

    def extractor_falso(imagen, ruc, nombre, tipo, cdcs):
        assert imagen[:2] == b"\xff\xd8"  # la IA recibe la página como JPG
        llamadas.append(cdcs)
        n = len(llamadas)
        if cdcs:  # página con QR: KuDE impreso y escaneado
            return {"contenido_pagina": "COMPROBANTES", "comprobantes": [_item(clase="ELECTRONICA", cdc="")]}
        if n == 1:  # dos facturas en la misma hoja, la segunda duplicada
            return {"contenido_pagina": "COMPROBANTES", "comprobantes": [_item(), _item(numero="001-001-0045646", total=60000)]}
        return {"contenido_pagina": "COMPROBANTES", "comprobantes": [_item()]}

    archivos = [
        ("lote.pdf", _pdf(["factura escaneada"])),
        ("kude.pdf", _pdf(["KuDE de factura electronica CDC " + " ".join(CDC[i:i + 4] for i in range(0, 44, 4))])),
        ("foto_qr.png", _png_con_qr(f"https://ekuatia.set.gov.py/consultas/qr?Id={CDC}")),
        ("repetida.jpg", _png_con_qr("no es un cdc")),
    ]
    resumen = procesador.procesar_lote(lote_id, archivos, extractor=extractor_falso)

    assert resumen["paginas"] == 4
    assert len(llamadas) == 4  # también el KuDE digital: la IA lee sus montos para Contadis
    comps = db.filas("SELECT * FROM comprobantes ORDER BY id")
    estados = [(c["clase"], c["estado"]) for c in comps]
    assert estados[0] == ("PREIMPRESA", "OK")
    assert estados[1] == ("PREIMPRESA", "REVISAR")  # total no cuadra
    assert ("ELECTRONICA", "EXCLUIDA") in estados
    assert comps[-1]["estado"] == "REVISAR" and any("duplicado" in e for e in comps[-1]["errores"])

    van, no_van = exportar.exportables(comps)
    assert len(van) == 1
    (nombre, contenido, cantidad), = exportar.generar_zips(van, "5850580", "04/2026", 1)
    assert nombre == "5850580_REG_042026_C0001.zip" and cantidad == 1
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        assert z.namelist() == ["5850580_REG_042026_C0001.txt"]
        linea = z.read(z.namelist()[0]).decode("utf-8").rstrip("\r\n")
    assert linea.split("\t") == ["2", "11", "1895550", "", "109", "09/04/2026", "18586180", "001-001-0045645",
                                 "51000", "0", "0", "51000", "1", "N", "S", "N", "N", "N", "", ""]
    assert exportar.generar_excel_control(van, no_van, {})[:2] == b"PK"


def test_division_en_varios_archivos(monkeypatch):
    monkeypatch.setattr(config, "MAX_FILAS_ARCHIVO", 2)
    c = {"tipo_comprobante": 109, "ruc": "1", "fecha": "01/04/2026", "timbrado": "12345678", "numero": "001-001-0000001",
         "gravada_10": 11, "gravada_5": 0, "exenta": 0, "total": 11, "condicion": 1, "imputa_iva": 1, "imputa_ire": 0,
         "imputa_irp": 0}
    zips = exportar.generar_zips([c] * 5, "80012345", "04/2026", 3)
    assert [z[0] for z in zips] == [f"80012345_REG_042026_C000{i}.zip" for i in (3, 4, 5)]
    assert [z[2] for z in zips] == [2, 2, 1]


# ---------- período por fecha de emisión ----------

def test_factura_de_otro_mes_va_a_su_periodo(base):
    cliente_id, lote_id = _cliente_y_lote()
    items = [_item(), _item(numero="001-001-0046021", fecha_emision="07/05/2026", gravada_10=85000, total=85000, iva_10=7727)]
    resumen = procesador.procesar_lote(lote_id, [("lote.pdf", _pdf(["x"]))],
                                       extractor=lambda *a: {"contenido_pagina": "COMPROBANTES", "comprobantes": items})
    assert resumen["OK"] == 2 and resumen["periodos"] == {"04/2026": 1, "05/2026": 1}
    mayo = db.fila("SELECT * FROM comprobantes WHERE numero='001-001-0046021'")
    assert mayo["periodo"] == "05/2026" and mayo["estado"] == "OK"
    assert "vino en el lote de 04/2026" in mayo["observaciones"]


# ---------- proveedores de IA ----------

from core import extractor  # noqa: E402


def test_normalizar_resultado():
    crudo = {"contenido_pagina": "raro", "comprobantes": [
        {"clase": "PREIMPRESA", "tipo_documento": "FACTURA", "total": "51.000", "gravada_10": 51000.0,
         "condicion": "contado", "campos_dudosos": None, "timbrado": 18586180}]}
    r = extractor.normalizar_resultado(crudo)
    c = r["comprobantes"][0]
    assert r["contenido_pagina"] == "COMPROBANTES"
    assert c["total"] == 51000 and c["gravada_10"] == 51000 and c["timbrado"] == "18586180"
    assert c["condicion"] == "NO_INDICA"  # valor fuera de la lista: se marca como no indicado
    assert "numero" in c["campos_dudosos"] and c["numero"] == ""  # campo que faltó: queda dudoso


def _imagen(tmp_path):
    from PIL import Image
    salida = io.BytesIO()
    Image.new("RGB", (50, 50), "white").save(salida, "JPEG")
    return salida.getvalue()


def test_extractor_gemini(tmp_path, monkeypatch):
    from google.genai import types
    capturado = {}

    class Respuesta:
        text = '{"contenido_pagina": "COMPROBANTES", "comprobantes": [{"clase": "PREIMPRESA", "total": 100}]}'

    class Modelos:
        def generate_content(self, **kwargs):
            capturado.update(kwargs)
            return Respuesta()

    class Cliente:
        models = Modelos()

    monkeypatch.setattr(config, "PROVEEDOR_IA", "gemini")
    monkeypatch.setattr(config, "GEMINI_MODELO", "gemini-prueba")
    monkeypatch.setattr(extractor, "_cliente_gemini", lambda: Cliente())
    r = extractor.extraer_pagina(_imagen(tmp_path), "5850580", "Cliente", "COMPRAS", [])
    assert r["comprobantes"][0]["total"] == 100
    assert capturado["model"] == "gemini-prueba"
    assert isinstance(capturado["contents"][0], types.Part)
    cfg = capturado["config"]
    assert cfg.response_mime_type == "application/json"
    assert "additionalProperties" not in json.dumps(cfg.response_json_schema)


def test_extractor_claude(tmp_path, monkeypatch):
    capturado = {}

    class Bloque:
        type = "text"
        text = '{"contenido_pagina": "NO_ES_COMPROBANTE", "comprobantes": []}'

    class Respuesta:
        stop_reason = "end_turn"
        content = [Bloque()]

    class Mensajes:
        def create(self, **kwargs):
            capturado.update(kwargs)
            return Respuesta()

    class Cliente:
        class beta:
            messages = Mensajes()

    monkeypatch.setattr(config, "PROVEEDOR_IA", "claude")
    monkeypatch.setattr(extractor, "_cliente_claude", lambda: Cliente())
    r = extractor.extraer_pagina(_imagen(tmp_path), "5850580", "Cliente", "COMPRAS", [])
    assert r == {"contenido_pagina": "NO_ES_COMPROBANTE", "comprobantes": []}
    assert capturado["model"] == config.MODELO and capturado["fallbacks"] == "default"


def test_problema_configuracion(monkeypatch):
    monkeypatch.setattr(config, "PROVEEDOR_IA", "gemini")
    monkeypatch.setattr(config, "GEMINI_MODELO", "")
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)
    assert "GEMINI_API_KEY" in config.problema_configuracion()
    monkeypatch.setenv("GEMINI_API_KEY", "x")
    assert "GEMINI_MODELO" in config.problema_configuracion()
    monkeypatch.setattr(config, "GEMINI_MODELO", "m")
    assert config.problema_configuracion() is None
    monkeypatch.setattr(config, "PROVEEDOR_IA", "otro")
    assert "claude" in config.problema_configuracion()


# ---------- reintentos ante el límite por minuto de la IA ----------

def test_reintenta_si_la_ia_pide_esperar(monkeypatch):
    from google.genai import errors
    monkeypatch.setattr(extractor.time, "sleep", lambda s: None)
    llamadas = []

    def llamar(imagen, pedido):
        llamadas.append(1)
        if len(llamadas) < 3:
            raise errors.ClientError(429, {"error": {"message": "Resource exhausted", "status": "RESOURCE_EXHAUSTED"}})
        return {"contenido_pagina": "NO_ES_COMPROBANTE", "comprobantes": []}

    assert extractor._con_reintentos(llamar, b"", "")["contenido_pagina"] == "NO_ES_COMPROBANTE"
    assert len(llamadas) == 3

    def falla(imagen, pedido):
        raise errors.ClientError(400, {"error": {"message": "API key not valid", "status": "INVALID_ARGUMENT"}})
    with pytest.raises(errors.ClientError):  # un error que no es de límite no se reintenta
        extractor._con_reintentos(falla, b"", "")


# ---------- usuarios ----------

def test_usuarios(base):
    assert not db.hay_usuarios()
    db.crear_usuario("Admin", "Administrador", "clave-segura-1", es_admin=True)
    u = db.autenticar("admin", "clave-segura-1")
    assert u and u["es_admin"] is True
    assert db.autenticar("admin", "otra") is None
    db.cambiar_clave(u["id"], "clave-nueva-2")
    assert db.autenticar("admin", "clave-nueva-2")


# ---------- imágenes temporales ----------

def _existe(base, ruta):
    return (base / ruta).exists()


def test_imagenes_se_borran_al_cerrar_comprobantes(base):
    cliente_id, lote_id = _cliente_y_lote()
    respuestas = [
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [_item()]},                                 # p1: OK
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [_item(numero="001-001-0000777", total=1)]},  # p2: REVISAR
        {"contenido_pagina": "CONTINUACION", "comprobantes": []},                                        # p3: sigue a p2
        {"contenido_pagina": "NO_ES_COMPROBANTE", "comprobantes": []},                                   # p4: hoja en blanco
    ]
    procesador.procesar_lote(lote_id, [("lote.pdf", _pdf(["1", "2", "3", "4"]))],
                             extractor=lambda *a: respuestas.pop(0))
    pag = {p["nro"]: p for p in db.filas("SELECT * FROM paginas ORDER BY nro")}
    rutas = {n: f"lote_{lote_id:06d}/pag_{n:04d}.jpg" for n in pag}

    # La hoja en blanco se borra enseguida; las demás se conservan para revisar.
    assert pag[4]["ruta"] is None and not _existe(base, rutas[4])
    assert all(_existe(base, rutas[n]) for n in (1, 2, 3))

    # Aprobar el comprobante de la p2 libera la p2 y su continuación (p3).
    c2 = db.fila("SELECT * FROM comprobantes WHERE pagina_id=%s", (pag[2]["id"],))
    assert pag[3]["comprobante_ref"] == c2["id"]
    db.actualizar_comprobante(c2["id"], {"estado": "APROBADA"}, "test", "aprobar")
    assert archivos.liberar([pag[2]["id"], pag[3]["id"]]) == 2
    assert not _existe(base, rutas[2]) and not _existe(base, rutas[3])

    # La p1 (OK) se conserva hasta que se exporta.
    assert archivos.liberar([pag[1]["id"]]) == 0 and _existe(base, rutas[1])
    db.ejecutar("UPDATE comprobantes SET exportado = true WHERE pagina_id=%s", (pag[1]["id"],))
    assert archivos.liberar([pag[1]["id"]]) == 1 and not _existe(base, rutas[1])
    assert archivos.bajar(None) is None


def test_pagina_con_dos_facturas_espera_a_ambas(base):
    cliente_id, lote_id = _cliente_y_lote()
    procesador.procesar_lote(lote_id, [("hoja.pdf", _pdf(["1"]))], extractor=lambda *a: {
        "contenido_pagina": "COMPROBANTES",
        "comprobantes": [_item(numero="001-001-0000001", total=1), _item(numero="001-001-0000002", total=1)]})
    pagina = db.fila("SELECT * FROM paginas")
    a, b = db.filas("SELECT id FROM comprobantes ORDER BY id")
    db.actualizar_comprobante(a["id"], {"estado": "APROBADA"}, "test", "aprobar")
    assert archivos.liberar([pagina["id"]]) == 0      # falta la otra factura de la hoja
    db.actualizar_comprobante(b["id"], {"estado": "DESCARTADA"}, "test", "descartar")
    assert archivos.liberar([pagina["id"]]) == 1


# ---------- Fase 2: ventas ----------

from core.validaciones import NUMERO_SIN_NOMBRE  # noqa: E402


@pytest.mark.parametrize("tipo_doc,numero,dv,esperado", [
    ("RUC", "80012345", "0", ("11", "80012345", "0")),
    ("RUC", "80.012.345-0", "", ("11", "80012345", "0")),
    ("CEDULA", "1.234.567", "", ("12", "1234567", "")),
    ("SIN_NOMBRE", "", "", ("15", NUMERO_SIN_NOMBRE, "")),
    ("RUC", "", "", ("15", NUMERO_SIN_NOMBRE, "")),          # factura sin datos del comprador
    ("PASAPORTE", "AB123456", "", ("13", "AB123456", "")),
])
def test_comprador_de_venta(tipo_doc, numero, dv, esperado):
    assert procesador.comprador_de_venta(tipo_doc, numero, dv) == esperado


def _venta(**cambios):
    c = _comprobante(ruc="80012345", dv="0", tipo_identificacion="11", ruc_receptor="5850580",
                     razon_social="Comprador SA")
    c.update(cambios)
    return c


def test_venta_correcta():
    assert validar_comprobante(_venta(), "5850580", "04/2026", tipo_lote="VENTAS") == []
    assert validar_comprobante(_venta(tipo_identificacion="12", ruc="1234567", dv=""), "5850580", "04/2026",
                               tipo_lote="VENTAS") == []
    assert validar_comprobante(_venta(tipo_identificacion="15", ruc=NUMERO_SIN_NOMBRE, dv="", razon_social=""),
                               "5850580", "04/2026", tipo_lote="VENTAS") == []


@pytest.mark.parametrize("cambio,texto", [
    ({"ruc_receptor": "1895550"}, "emitida por el RUC 1895550"),
    ({"dv": "1"}, "DV del RUC del comprador incorrecto"),
    ({"tipo_identificacion": "15", "ruc": "0", "dv": ""}, NUMERO_SIN_NOMBRE),
    ({"tipo_identificacion": "12", "ruc": "1.234.567", "dv": ""}, "sin puntos"),
    ({"tipo_identificacion": "13", "ruc": "AB1", "dv": "", "razon_social": ""}, "Falta el nombre del comprador"),
    ({"tipo_identificacion": "99"}, "Tipo de identificación"),
])
def test_venta_con_errores(cambio, texto):
    errores = validar_comprobante(_venta(**cambio), "5850580", "04/2026", tipo_lote="VENTAS")
    assert any(texto in e for e in errores), errores


def _fila_base(**cambios):
    c = {"tipo_comprobante": 109, "tipo_identificacion": "11", "ruc": "80012345", "razon_social": "Comprador SA",
         "fecha": "09/04/2026", "timbrado": "12345678", "numero": "001-001-0000010", "gravada_10": 11000,
         "gravada_5": 2100, "exenta": 500, "total": 13600, "condicion": 2, "imputa_iva": 1, "imputa_ire": 0,
         "imputa_irp": 0}
    c.update(cambios)
    return c


def test_fila_venta_19_campos():
    f = exportar.fila_venta(_fila_base())
    assert len(f) == 19 and len(exportar.COLUMNAS_VENTAS) == 19
    assert f == ["1", "11", "80012345", "", "109", "09/04/2026", "12345678", "001-001-0000010",
                 "11000", "2100", "500", "13600", "2", "N", "S", "N", "N", "", ""]
    # Pasaporte: la razón social sí va. Ticket: en ventas conserva los montos.
    assert exportar.fila_venta(_fila_base(tipo_identificacion="13", ruc="AB123"))[3] == "Comprador SA"
    ticket = exportar.fila_venta(_fila_base(tipo_comprobante=112, tipo_identificacion="15", ruc=NUMERO_SIN_NOMBRE))
    assert ticket[1:4] == ["15", NUMERO_SIN_NOMBRE, ""] and ticket[8:12] == ["11000", "2100", "500", "13600"]
    assert len(exportar.fila_compra(_fila_base())) == 20


def test_lote_de_ventas_completo(base):
    cliente_id, _ = _cliente_y_lote()
    lote_id = db.ejecutar("INSERT INTO lotes (cliente_id, tipo, periodo, creado_por) "
                          "VALUES (%s, 'VENTAS', '04/2026', 'test') RETURNING id", (cliente_id,))
    emitida = dict(ruc_emisor="5850580", dv_emisor="6", razon_social_emisor="Cliente Prueba",
                   timbrado="12345678", gravada_10=11000, total=11000, iva_10=1000)
    respuestas = [
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [
            _item(**emitida, numero="001-001-0000010", ruc_receptor="80012345", dv_receptor="0",
                  tipo_doc_receptor="RUC", nombre_receptor="Comprador SA"),
            _item(**emitida, numero="001-001-0000011", ruc_receptor="", dv_receptor="",
                  tipo_doc_receptor="SIN_NOMBRE", nombre_receptor="")]},
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [
            _item(**emitida, numero="001-001-0000012", ruc_receptor="3.456.789", dv_receptor="",
                  tipo_doc_receptor="CEDULA", nombre_receptor="Juan Pérez")]},
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [   # factura de otro emisor: no es venta del cliente
            _item(numero="001-001-0000099", tipo_doc_receptor="RUC", nombre_receptor="X")]},
    ]
    resumen = procesador.procesar_lote(lote_id, [("ventas.pdf", _pdf(["1", "2", "3"]))],
                                       extractor=lambda *a: respuestas.pop(0))
    assert resumen["OK"] == 3 and resumen["REVISAR"] == 1
    malo = db.fila("SELECT * FROM comprobantes WHERE estado='REVISAR'")
    assert any("emitida por el RUC" in e for e in malo["errores"])

    comps = db.filas("SELECT * FROM comprobantes WHERE tipo_lote='VENTAS' ORDER BY id")
    van, _ = exportar.exportables(comps)
    (nombre, contenido, cantidad), = exportar.generar_zips(van, "5850580", "04/2026", 1, "VENTAS")
    assert nombre == "5850580_REG_042026_V0001.zip" and cantidad == 3
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        lineas = [l.split("\t") for l in z.read(z.namelist()[0]).decode().removesuffix("\r\n").split("\r\n")]
    assert [l[:4] for l in lineas] == [["1", "11", "80012345", ""], ["1", "15", NUMERO_SIN_NOMBRE, ""],
                                       ["1", "12", "3456789", ""]]
    assert all(len(l) == 19 for l in lineas)

    # En ventas, el duplicado es el mismo timbrado y número aunque cambie el comprador.
    assert db.es_duplicado(cliente_id, "VENTAS", "999", "12345678", "001-001-0000010")
    assert not db.es_duplicado(cliente_id, "COMPRAS", "999", "12345678", "001-001-0000010")


# ---------- resumen para Contadis ----------

from core import contadis  # noqa: E402


@pytest.mark.parametrize("con_iva,iva_impreso,divisor,esperado", [
    (1_000_000, 90_909, 11, (909_091, 90_909)),      # captura de compras de Contadis
    (1_100_000, 100_000, 11, (1_000_000, 100_000)),  # captura de ventas de Contadis
    (51_000, 4_636, 11, (46_364, 4_636)),            # factura de Maxi Copias
    (51_000, 0, 11, (46_364, 4_636)),                # sin IVA impreso: se calcula
    (51_000, 9_999, 11, (46_364, 4_636)),            # IVA impreso incoherente: se usa el calculado
    (21_000, 1_000, 21, (20_000, 1_000)),            # 5%
])
def test_separar_iva(con_iva, iva_impreso, divisor, esperado):
    assert contadis.separar_iva(con_iva, iva_impreso, divisor) == esperado


def _para_contadis(**cambios):
    c = {"estado": "OK", "clase": "PREIMPRESA", "tipo_comprobante": 109, "numero": "001-001-0045645",
         "fecha": "09/04/2026", "condicion": 1, "timbrado": "18586180", "ruc": "1895550", "dv": "9",
         "razon_social": "Maxi Copias", "tipo_identificacion": "11", "gravada_10": 51000, "gravada_5": 0,
         "exenta": 0, "total": 51000, "iva_10": 4636, "iva_5": 0, "concepto": "Útiles de librería",
         "errores": [], "pagina_id": None}
    c.update(cambios)
    return c


def test_fila_contadis_compras_y_ventas():
    valores, avisos = contadis.fila(_para_contadis(), "COMPRAS")
    assert dict(zip(contadis.COLUMNAS["COMPRAS"], valores)) | {"Fecha": None} == {
        "Est.": "001", "Pto.": "001", "Número": "0045645", "Comprobante": "FACTURA", "Razón social": "Maxi Copias",
        "Fecha": None, "Condición": "Contado", "Timbrado": "18586180", "RUC": "1895550-9",
        "Gravadas 5%": 0, "Gravadas 10%": 46364, "I.V.A. 5%": 0, "I.V.A. 10%": 4636, "Exenta": 0, "Total": 51000,
        "Concepto": "Útiles de librería", "Electrónica": "", "Cargado": "", "Observaciones": ""}
    assert avisos == []
    sin_nombre = dict(zip(contadis.COLUMNAS["VENTAS"], contadis.fila(
        _para_contadis(tipo_identificacion="15", ruc=NUMERO_SIN_NOMBRE, dv="", razon_social=""), "VENTAS")[0]))
    assert sin_nombre["Razón social"] == "CLIENTE OCASIONAL" and sin_nombre["RUC / CI"] == NUMERO_SIN_NOMBRE
    cedula = dict(zip(contadis.COLUMNAS["VENTAS"], contadis.fila(
        _para_contadis(tipo_identificacion="12", ruc="3456789", dv="", razon_social="Juan Pérez"), "VENTAS")[0]))
    assert cedula["RUC / CI"] == "3456789"


def test_resumen_contadis_excel():
    from openpyxl import load_workbook
    compras = [_para_contadis(numero="001-001-0000002", fecha="20/04/2026"),
               _para_contadis(numero="001-001-0000001", fecha="05/04/2026", clase="ELECTRONICA", estado="EXCLUIDA"),
               _para_contadis(numero="001-001-0000003", estado="REVISAR", errores=["DV incorrecto"]),
               _para_contadis(numero="001-001-0000004", estado="DESCARTADA"),
               _para_contadis(numero="001-001-0000005", clase="ELECTRONICA", estado="EXCLUIDA", total=0,
                              gravada_10=0, iva_10=0)]
    ventas = [_para_contadis(tipo_identificacion="15", ruc=NUMERO_SIN_NOMBRE, dv="", razon_social="")]
    datos = contadis.generar_resumen(compras, ventas, {"nombre": "Cliente", "ruc": "5850580", "dv": "6"}, "04/2026")
    wb = load_workbook(io.BytesIO(datos))
    assert wb.sheetnames == ["Resumen", "Compras", "Ventas", "Pendientes"]
    ws = wb["Compras"]
    assert [ws.cell(r, 3).value for r in (2, 3)] == ["0000001", "0000002"]   # ordenadas por fecha
    assert ws.cell(2, 17).value == "Sí"                                       # electrónica incluida
    assert ws.cell(5, 1).value == "TOTALES" and ws.cell(5, 15).value == "=SUM(O2:O3)"
    pendientes = [(r[0].value, r[2].value) for r in wb["Pendientes"].iter_rows(min_row=2)]
    assert pendientes == [("COMPRAS", "001-001-0000003"), ("COMPRAS", "001-001-0000005")]
    assert wb["Ventas"].cell(2, 5).value == "CLIENTE OCASIONAL"
    resumen = [[c.value for c in fila] for fila in wb["Resumen"].iter_rows(min_row=6, max_row=7)]
    assert resumen[0][:3] == ["Compras", 2, 2] and resumen[0][-1] == 102000
    assert resumen[1][:3] == ["Ventas", 1, 0]


def test_electronica_guarda_montos_para_contadis(base):
    cliente_id, lote_id = _cliente_y_lote()
    kude = "KuDE CDC " + " ".join(CDC[i:i + 4] for i in range(0, 44, 4))
    respuestas = [
        {"contenido_pagina": "COMPROBANTES", "comprobantes": [
            _item(clase="ELECTRONICA", cdc="", ruc_emisor="80012345", dv_emisor="0", fecha_emision="15/09/2026",
                  numero="001-002-0000123", gravada_10=110000, total=110000, iva_10=10000, concepto="Combustible")]},
    ]

    def extractor(imagen, ruc, nombre, tipo, cdcs):
        assert cdcs == [CDC]  # la IA recibe el CDC leído del texto
        return respuestas.pop(0)

    procesador.procesar_lote(lote_id, [("kude.pdf", _pdf([kude]))], extractor=extractor)
    c = db.fila("SELECT * FROM comprobantes")
    assert (c["clase"], c["estado"], c["cdc"]) == ("ELECTRONICA", "EXCLUIDA", CDC)
    assert (c["total"], c["iva_10"], c["concepto"]) == (110000, 10000, "Combustible")
    assert (c["ruc"], c["numero"], c["fecha"], c["periodo"]) == ("80012345", "001-002-0000123", "15/09/2026", "09/2026")
    cargar, pendientes = contadis.separar([c])
    assert len(cargar) == 1 and not pendientes



def test_electronica_sin_montos_si_falla_la_ia(base):
    cliente_id, lote_id = _cliente_y_lote()
    kude = "KuDE CDC " + " ".join(CDC[i:i + 4] for i in range(0, 44, 4))

    def falla(*a):
        raise RuntimeError("sin conexión")

    procesador.procesar_lote(lote_id, [("kude.pdf", _pdf([kude]))], extractor=falla)
    c = db.fila("SELECT * FROM comprobantes")
    assert (c["clase"], c["estado"], c["total"]) == ("ELECTRONICA", "EXCLUIDA", 0)
    assert "Sin montos" in c["observaciones"]
    cargar, pendientes = contadis.separar([c])
    assert not cargar and len(pendientes) == 1
